"""Remesas SEPA para gestorías: deudores, remesas, facturas automáticas, exportación Excel/CSV y detección de cobros."""
import io
import csv
import uuid
from datetime import datetime, timezone
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import Response
from pydantic import BaseModel
from bson import ObjectId

from database import db
from auth import get_current_user


async def require_gestoria(user=Depends(get_current_user)):
    """Remesas solo para gestorías, planes Multiempresas y administradores; cada uno ve solo las suyas."""
    if user.get("role") in ("gestoria", "admin") or str(user.get("plan") or "").startswith("multiempresas"):
        return user
    raise HTTPException(status_code=403, detail="Las remesas SEPA están disponibles para gestorías y planes Multiempresas.")

rem = APIRouter(prefix="/api/gestoria/remesas", tags=["remesas"])

COLUMNS = ["Nº", "Importe", "Fecha cobro", "Nombre deudor", "Identificador deudor", "IBAN", "Mandato", "Fecha mandato", "Referencia operación", "Concepto",
           "Base imponible", "IVA %", "Cuota IVA", "Exento", "Factura"]
STATUSES = {"borrador", "emitida", "enviada", "cobrada", "parcial"}
OPEN = ["borrador", "emitida", "enviada", "parcial"]
DEFAULT_EXEMPT = "Exento art. 20.Uno Ley 37/1992 del IVA"


def _now():
    return datetime.now(timezone.utc).isoformat()


def _r(x):
    return round(float(x or 0), 2)


def _num(v) -> float:
    s = str(v if v is not None else "").strip().replace("€", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return 0.0


def iban_valid(iban: str) -> bool:
    s = (iban or "").replace(" ", "").upper()
    if len(s) < 15 or len(s) > 34 or not s[:2].isalpha() or not s[2:4].isdigit():
        return False
    if s.startswith("ES") and len(s) != 24:
        return False
    num = "".join(str(int(c, 36)) for c in s[4:] + s[:4])
    return int(num) % 97 == 1


def line_warnings(ln: dict) -> list:
    w = []
    if (ln.get("nif") or "").strip() and not __import__("spanish_tax").validate_nif(ln["nif"]):
        w.append("NIF no válido")
    if _r(ln.get("amount")) <= 0:
        w.append("Importe 0")
    if not iban_valid(ln.get("iban", "")):
        w.append("IBAN no válido")
    if not (ln.get("mandate") or "").strip():
        w.append("Falta mandato")
    if not (ln.get("mandate_date") or "").strip():
        w.append("Falta fecha de mandato")
    if not (ln.get("name") or "").strip():
        w.append("Falta nombre del deudor")
    return w


class DebtorInput(BaseModel):
    name: str = ""
    nif: str = ""
    iban: str = ""
    mandate: str = ""
    mandate_date: str = ""
    email: str = ""
    address: str = ""
    amount: float = 0.0
    iva_rate: float = 21
    vat_exempt: bool = False
    exempt_reason: str = ""
    concept: str = ""
    reference: str = ""
    client_user_id: str = ""
    active: bool = True


class LineInput(BaseModel):
    id: str = ""
    debtor_id: str = ""
    name: str = ""
    nif: str = ""
    iban: str = ""
    mandate: str = ""
    mandate_date: str = ""
    email: str = ""
    amount: float = 0.0
    iva_rate: float = 21
    vat_exempt: bool = False
    exempt_reason: str = ""
    reference: str = ""
    concept: str = ""
    invoice_id: str = ""
    invoice_number: str = ""
    paid: bool = False
    paid_date: str = ""


class RemesaInput(BaseModel):
    name: str
    charge_date: str = ""
    reference: str = ""
    lines: List[LineInput] = []


class CreateInput(BaseModel):
    name: str
    charge_date: str = ""
    reference: str = ""
    from_debtors: bool = True


def _debtor_doc(d: DebtorInput) -> dict:
    x = d.model_dump()
    x["iban"] = x["iban"].replace(" ", "").upper()
    x["nif"] = x["nif"].strip().upper()
    return x


def _line_from_debtor(d: dict) -> dict:
    return {"id": str(uuid.uuid4()), "debtor_id": d["id"], "name": d.get("name", ""), "nif": d.get("nif", ""), "iban": d.get("iban", ""),
            "mandate": d.get("mandate", ""), "mandate_date": d.get("mandate_date", ""), "email": d.get("email", ""),
            "amount": _r(d.get("amount")), "iva_rate": d.get("iva_rate", 21), "vat_exempt": bool(d.get("vat_exempt")),
            "exempt_reason": d.get("exempt_reason", ""), "reference": d.get("reference", ""),
            "concept": d.get("concept", ""), "invoice_id": "", "invoice_number": "", "paid": False, "paid_date": ""}


def _base_from_total(total: float, rate: float) -> float:
    base = round(total / (1 + rate / 100), 2)
    for delta in (0, 0.01, -0.01, 0.02, -0.02):
        b = round(base + delta, 2)
        if round(b + round(b * rate / 100, 2), 2) == round(total, 2):
            return b
    return base


def _rate(ln: dict) -> float:
    return 0.0 if ln.get("vat_exempt") else float(ln.get("iva_rate") or 0)


def _tax(ln: dict):
    """El importe del recibo es el TOTAL con IVA incluido: devuelve (base, cuota IVA)."""
    total = _r(ln.get("amount"))
    base = _base_from_total(total, _rate(ln))
    return base, _r(total - base)


def _summary(r: dict) -> dict:
    lines = r.get("lines", [])
    for ln in lines:
        ln["warnings"] = line_warnings(ln)
        ln["base"], ln["iva"] = _tax(ln)
    r["total"] = _r(sum(_r(ln.get("amount")) for ln in lines))
    r["base_total"] = _r(sum(ln["base"] for ln in lines))
    r["iva_total"] = _r(r["total"] - r["base_total"])
    r["count"] = len(lines)
    r["paid_total"] = _r(sum(_r(ln.get("amount")) for ln in lines if ln.get("paid")))
    r["pending_total"] = _r(r["total"] - r["paid_total"])
    r["pending_count"] = sum(1 for ln in lines if not ln.get("paid"))
    r["invoiced"] = sum(1 for ln in lines if ln.get("invoice_id"))
    r["warnings"] = sum(1 for ln in lines if ln["warnings"])
    return r


async def _get(rid, g):
    r = await db.remesas.find_one({"id": rid, "gestoria_id": g["id"]}, {"_id": 0})
    if not r:
        raise HTTPException(status_code=404, detail="Remesa no encontrada")
    return r


# ---------- Deudores ----------
@rem.get("/deudores")
async def list_debtors(g=Depends(require_gestoria)):
    rows = await db.remesa_debtors.find({"gestoria_id": g["id"]}, {"_id": 0}).sort("name", 1).to_list(5000)
    for d in rows:
        d["warnings"] = line_warnings(d)
    return rows


@rem.post("/deudores")
async def create_debtor(data: DebtorInput, g=Depends(require_gestoria)):
    doc = {**_debtor_doc(data), "id": str(uuid.uuid4()), "gestoria_id": g["id"], "created_at": _now()}
    await db.remesa_debtors.insert_one(doc)
    doc.pop("_id", None)
    return doc


@rem.put("/deudores/{debtor_id}")
async def update_debtor(debtor_id: str, data: DebtorInput, g=Depends(require_gestoria)):
    r = await db.remesa_debtors.update_one({"id": debtor_id, "gestoria_id": g["id"]}, {"$set": _debtor_doc(data)})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Deudor no encontrado")
    return {"status": "ok"}


@rem.delete("/deudores/{debtor_id}")
async def delete_debtor(debtor_id: str, g=Depends(require_gestoria)):
    await db.remesa_debtors.delete_one({"id": debtor_id, "gestoria_id": g["id"]})
    return {"status": "ok"}


@rem.post("/deudores/desde-clientes")
async def debtors_from_clients(g=Depends(require_gestoria)):
    n = 0
    if g.get("role") != "gestoria":
        cid = await __import__("server").active_cid(g)
        async for c in db.contacts.find({"user_id": g["id"], "company_id": cid, "kind": {"$ne": "provider"}}, {"_id": 0}):
            if await db.remesa_debtors.find_one({"gestoria_id": g["id"], "contact_id": c.get("id")}):
                continue
            await db.remesa_debtors.insert_one({**DebtorInput(name=c.get("name", ""), nif=(c.get("nif") or "").upper(), email=c.get("email", ""),
                                                              address=c.get("address", "")).model_dump(),
                                                "contact_id": c.get("id"), "id": str(uuid.uuid4()), "gestoria_id": g["id"], "created_at": _now()})
            n += 1
        return {"created": n}
    async for u in db.users.find({"gestoria_id": g["id"]}):
        uid = str(u["_id"])
        if await db.remesa_debtors.find_one({"gestoria_id": g["id"], "client_user_id": uid}):
            continue
        comp = await db.companies.find_one({"user_id": uid}) or {}
        name = comp.get("name") or " ".join(x for x in [u.get("name", ""), u.get("last_name", "")] if x)
        await db.remesa_debtors.insert_one({**DebtorInput(name=name, nif=(comp.get("nif") or u.get("tax_id") or "").upper(),
                                                          email=u.get("email", ""), address=comp.get("address", ""), client_user_id=uid,
                                                          concept="Asesoría mes vista").model_dump(),
                                            "id": str(uuid.uuid4()), "gestoria_id": g["id"], "created_at": _now()})
        n += 1
    return {"created": n}


# ---------- Remesas ----------
@rem.get("")
async def list_remesas(g=Depends(require_gestoria)):
    rows = await db.remesas.find({"gestoria_id": g["id"]}, {"_id": 0}).sort("created_at", -1).to_list(1000)
    out = []
    for r in rows:
        s = _summary(r)
        s.pop("lines", None)
        out.append(s)
    return out


@rem.post("")
async def create_remesa(data: CreateInput, g=Depends(require_gestoria)):
    lines = []
    if data.from_debtors:
        async for d in db.remesa_debtors.find({"gestoria_id": g["id"], "active": True}, {"_id": 0}).sort("name", 1):
            lines.append(_line_from_debtor(d))
    doc = {"id": str(uuid.uuid4()), "gestoria_id": g["id"], "name": data.name, "charge_date": data.charge_date,
           "reference": data.reference, "status": "borrador", "lines": lines, "created_at": _now()}
    await db.remesas.insert_one(doc)
    doc.pop("_id", None)
    return _summary(doc)


@rem.get("/{rid}")
async def get_remesa(rid: str, g=Depends(require_gestoria)):
    return _summary(await _get(rid, g))


@rem.put("/{rid}")
async def update_remesa(rid: str, data: RemesaInput, g=Depends(require_gestoria)):
    r = await _get(rid, g)
    old = {ln["id"]: ln for ln in r.get("lines", [])}
    lines = []
    for ln in data.lines:
        x = ln.model_dump()
        x["id"] = x["id"] or str(uuid.uuid4())
        x["iban"] = x["iban"].replace(" ", "").upper()
        x["nif"] = x["nif"].strip().upper()
        prev = old.get(x["id"])
        if prev and prev.get("invoice_id"):
            x["invoice_id"], x["invoice_number"] = prev["invoice_id"], prev.get("invoice_number", "")
            x["issue_date"] = prev.get("issue_date", "")
            if _r(prev.get("amount")) != _r(x["amount"]) or _rate(prev) != _rate(x):
                raise HTTPException(status_code=400, detail=f"La línea de {x['name']} ya tiene factura ({x['invoice_number']}); no se puede cambiar su importe ni su IVA. Emite una rectificativa.")
        lines.append(x)
        if x.get("debtor_id"):
            upd = {k: x[k] for k in ("name", "nif", "iban", "mandate", "mandate_date", "email", "concept", "reference") if x.get(k)}
            upd.update({"iva_rate": x["iva_rate"], "vat_exempt": x["vat_exempt"], "exempt_reason": x["exempt_reason"]})
            await db.remesa_debtors.update_one({"id": x["debtor_id"], "gestoria_id": g["id"]}, {"$set": upd})
    await db.remesas.update_one({"id": rid}, {"$set": {"name": data.name, "charge_date": data.charge_date, "reference": data.reference,
                                                        "lines": lines, "updated_at": _now()}})
    return _summary(await _get(rid, g))


@rem.delete("/{rid}")
async def delete_remesa(rid: str, g=Depends(require_gestoria)):
    r = await _get(rid, g)
    if any(ln.get("invoice_id") for ln in r.get("lines", [])):
        raise HTTPException(status_code=400, detail="La remesa ya tiene facturas emitidas; no se puede eliminar.")
    await db.remesas.delete_one({"id": rid})
    return {"status": "ok"}


class DuplicateInput(BaseModel):
    name: str
    charge_date: str = ""


@rem.post("/{rid}/duplicar")
async def duplicate(rid: str, data: DuplicateInput, g=Depends(require_gestoria)):
    r = await _get(rid, g)
    lines = [{**ln, "id": str(uuid.uuid4()), "invoice_id": "", "invoice_number": "", "issue_date": "", "paid": False, "paid_date": ""} for ln in r.get("lines", [])]
    for ln in lines:
        ln.pop("warnings", None)
    doc = {"id": str(uuid.uuid4()), "gestoria_id": g["id"], "name": data.name, "charge_date": data.charge_date,
           "reference": r.get("reference", ""), "status": "borrador", "lines": lines, "created_at": _now()}
    await db.remesas.insert_one(doc)
    doc.pop("_id", None)
    return _summary(doc)


class StatusInput(BaseModel):
    status: str


@rem.post("/{rid}/estado")
async def set_status(rid: str, data: StatusInput, g=Depends(require_gestoria)):
    if data.status not in STATUSES:
        raise HTTPException(status_code=400, detail="Estado no válido")
    await _get(rid, g)
    await db.remesas.update_one({"id": rid}, {"$set": {"status": data.status}})
    return {"status": data.status}


class PaidInput(BaseModel):
    paid: bool
    date: str = ""


async def _mark_line(r, ln, paid, pdate, source="manual"):
    ln["paid"], ln["paid_date"] = paid, (pdate or datetime.now(timezone.utc).date().isoformat())[:10] if paid else ""
    if ln.get("invoice_id"):
        st = {"status": "paid", "payment": {"status": "paid", "method": "sepa", "date": ln["paid_date"], "source": source, "amount": _r(ln["amount"])}} \
            if paid else {"status": "pending", "payment": None}
        await db.invoices.update_one({"id": ln["invoice_id"]}, {"$set": st})


async def _save_lines(r):
    lines = r["lines"]
    for ln in lines:
        for k in ("warnings", "base", "iva"):
            ln.pop(k, None)
    allpaid = lines and all(ln.get("paid") for ln in lines)
    anypaid = any(ln.get("paid") for ln in lines)
    status = "cobrada" if allpaid else ("parcial" if anypaid else (r.get("status") if r.get("status") in ("borrador", "emitida", "enviada") else "enviada"))
    await db.remesas.update_one({"id": r["id"]}, {"$set": {"lines": lines, "status": status}})
    return status


@rem.post("/{rid}/lineas/{line_id}/cobro")
async def mark_paid(rid: str, line_id: str, data: PaidInput, g=Depends(require_gestoria)):
    r = await _get(rid, g)
    ln = next((x for x in r["lines"] if x["id"] == line_id), None)
    if not ln:
        raise HTTPException(status_code=404, detail="Línea no encontrada")
    await _mark_line(r, ln, data.paid, data.date)
    await _save_lines(r)
    return _summary(await _get(rid, g))


# ---------- Datos fiscales del emisor (gestoría) ----------
FISCAL_FIELDS = ("name", "legal_name", "nif", "address", "email", "phone", "tax_type", "legal_notice", "invoice_footer", "sepa_creditor_id", "sepa_iban", "sepa_bic")
REQUIRED = {"name": "Nombre comercial o razón social", "nif": "NIF/CIF", "address": "Domicilio fiscal completo"}


class FiscalInput(BaseModel):
    name: str = ""
    legal_name: str = ""
    nif: str = ""
    address: str = ""
    email: str = ""
    phone: str = ""
    tax_type: str = "sociedad"
    legal_notice: str = ""
    invoice_footer: str = ""
    sepa_creditor_id: str = ""
    sepa_iban: str = ""
    sepa_bic: str = ""


def _sepa_missing(c: dict) -> list:
    miss = []
    if not iban_valid(c.get("sepa_iban", "")):
        miss.append("IBAN de abono de la gestoría")
    if not (c.get("sepa_creditor_id") or (c.get("nif") or "").strip()):
        miss.append("Identificador de acreedor SEPA")
    return miss


def _issuer_out(c: dict) -> dict:
    from remesa_docs import creditor_id_from_nif
    return {**{k: c.get(k, "") for k in FISCAL_FIELDS}, "id": c.get("id"), "missing": _missing(c), "sepa_missing": _sepa_missing(c),
            "suggested_creditor_id": creditor_id_from_nif(c.get("nif", ""))}


def _missing(c: dict) -> list:
    import spanish_tax
    miss = [label for k, label in REQUIRED.items() if not (c.get(k) or "").strip()]
    if (c.get("nif") or "").strip() and not spanish_tax.validate_nif(c["nif"]):
        miss.append("NIF/CIF válido")
    return miss


@rem.get("/ajustes/emisor")
async def get_issuer(g=Depends(require_gestoria)):
    c = await __import__("server").active_company(g)
    return _issuer_out(c)


@rem.put("/ajustes/emisor")
async def set_issuer(data: FiscalInput, g=Depends(require_gestoria)):
    c = await __import__("server").active_company(g)
    upper = ("nif", "sepa_creditor_id", "sepa_iban", "sepa_bic")
    upd = {k: (v.replace(" ", "").upper() if k in upper else v.strip()) for k, v in data.model_dump().items()}
    if upd["tax_type"] not in ("autonomo", "sociedad", "empresa"):
        upd["tax_type"] = "sociedad"
    if upd["sepa_iban"] and not iban_valid(upd["sepa_iban"]):
        raise HTTPException(status_code=400, detail="El IBAN de abono de la gestoría no es válido.")
    await db.companies.update_one({"id": c["id"]}, {"$set": upd})
    c.update(upd)
    return _issuer_out(c)


# ---------- Facturas ----------

@rem.post("/{rid}/facturas")
async def create_invoices(rid: str, g=Depends(require_gestoria)):
    srv = __import__("server")
    r = await _get(rid, g)
    company = await srv.active_company(g)
    miss = _missing(company)
    if miss:
        raise HTTPException(status_code=400, detail="DATOS_FISCALES: Completa los datos fiscales de la gestoría antes de emitir facturas: " + ", ".join(miss) + ".")
    import spanish_tax
    issue = datetime.now(timezone.utc).date().isoformat()
    created, errors, skipped = 0, [], 0
    for i, ln in enumerate(r["lines"], 1):
        who = (ln.get("name") or "").strip() or f"Recibo nº {i}"
        if ln.get("invoice_id"):
            skipped += 1
            continue
        rate = _rate(ln)
        probs = []
        if not (ln.get("name") or "").strip():
            probs.append("falta el nombre del cliente")
        if _r(ln.get("amount")) <= 0:
            probs.append("el importe es 0")
        nif = (ln.get("nif") or "").strip().upper()
        if not nif:
            probs.append("falta el NIF/CIF del cliente")
        elif not spanish_tax.validate_nif(nif):
            probs.append(f"el NIF/CIF {nif} no es válido")
        if rate not in (0, 4, 10, 21):
            probs.append(f"el IVA {rate:g}% no es válido (0, 4, 10 o 21)")
        if probs:
            errors.append(f"{who}: " + ", ".join(probs))
            continue
        base, _ = _tax(ln)
        reason = (ln.get("exempt_reason") or DEFAULT_EXEMPT).strip()
        notes = f"Se cargará en su cuenta {ln.get('iban', '')} (mandato {ln.get('mandate', '')})."
        if ln.get("vat_exempt"):
            notes = f"Operación exenta de IVA: {reason}. " + notes
        try:
            inv = await srv._make_invoice(g, company, srv.InvoiceInput(
                issue_date=issue,
                client=srv.Client(name=ln["name"].strip(), nif=nif, email=ln.get("email", "")),
                line_items=[srv.LineItem(description=ln.get("concept") or ln.get("reference") or r["name"],
                                         detail=f"Remesa {r['name']} · Mandato {ln.get('mandate', '')}", quantity=1, unit_price=base,
                                         iva_rate=rate, iva_type="exento" if (ln.get("vat_exempt") or not rate) else "general")],
                iva_rate=rate, status="pending", payment_method="Domiciliación bancaria SEPA", iban=ln.get("iban", ""),
                due_date=r.get("charge_date") or issue, notes=notes))
            if abs(_r(inv["total"]) - _r(ln["amount"])) >= 0.01:
                errors.append(f"{who}: la factura {inv['number']} suma {inv['total']:.2f} € y el recibo {_r(ln['amount']):.2f} €")
            ln["invoice_id"], ln["invoice_number"], ln["issue_date"] = inv["id"], inv["number"], issue
            await db.invoices.update_one({"id": inv["id"]}, {"$set": {"remesa_id": rid, "remesa_line_id": ln["id"]}})
            created += 1
        except HTTPException as e:
            errors.append(f"{who}: {e.detail}")
        except Exception as e:
            errors.append(f"{who}: error al crear la factura ({e})")
    for ln in r["lines"]:
        for k in ("warnings", "base", "iva"):
            ln.pop(k, None)
    upd = {"lines": r["lines"]}
    if created and r.get("status") == "borrador":
        upd["status"] = "emitida"
    await db.remesas.update_one({"id": rid}, {"$set": upd})
    return {"created": created, "skipped": skipped, "errors": errors, "remesa": _summary(await _get(rid, g))}


# ---------- Recibos PDF y fichero SEPA ----------
def _doc_lines(r: dict, line_id: str = "") -> list:
    s = _summary(r)
    out = []
    for i, ln in enumerate(s["lines"], 1):
        if not line_id or ln["id"] == line_id:
            out.append({**ln, "_idx": i})
    return out


@rem.get("/{rid}/recibos")
async def receipts_pdf(rid: str, line_id: str = "", g=Depends(require_gestoria)):
    from remesa_docs import build_receipts_pdf
    r = await _get(rid, g)
    lines = [ln for ln in _doc_lines(r, line_id) if _r(ln.get("amount")) > 0]
    if not lines:
        raise HTTPException(status_code=404, detail="No hay recibos con importe en esta remesa.")
    pdf = build_receipts_pdf(r, lines, await __import__("server").active_company(g))
    safe = "".join(c if c.isalnum() else "_" for c in (lines[0]["name"] if line_id else r["name"]))[:60] or "recibos"
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="recibo_{safe}.pdf"'})


@rem.get("/{rid}/sepa-xml")
async def sepa_xml(rid: str, g=Depends(require_gestoria)):
    from remesa_docs import build_sepa_xml
    r = await _get(rid, g)
    company = await __import__("server").active_company(g)
    miss = _missing(company) + _sepa_missing(company)
    if miss:
        raise HTTPException(status_code=400, detail="DATOS_FISCALES: Para generar el fichero SEPA completa: " + ", ".join(miss) + ".")
    lines = [ln for ln in _doc_lines(r) if not ln.get("paid")]
    bad = [f"{ln.get('name') or 'Recibo nº ' + str(ln['_idx'])}: {', '.join(w for w in ln['warnings'] if w != 'NIF no válido')}"
           for ln in lines if any(w != "NIF no válido" for w in ln["warnings"])]
    if bad:
        raise HTTPException(status_code=400, detail="Corrige estos recibos antes de generar el XML: " + " · ".join(bad))
    if not lines:
        raise HTTPException(status_code=400, detail="No hay recibos pendientes de cobro en esta remesa.")
    xml = build_sepa_xml(r, lines, company)
    safe = "".join(c if c.isalnum() else "_" for c in r["name"])[:60] or "remesa"
    return Response(xml, media_type="application/xml", headers={"Content-Disposition": f'attachment; filename="{safe}_pain008.xml"'})


# ---------- Importar / Exportar ----------
def _rows_from_file(name: str, raw: bytes) -> list:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), data_only=True).active
        rows = [[("" if c is None else (c.date().isoformat() if hasattr(c, "date") and not isinstance(c, str) else str(c))) for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        text = raw.decode("utf-8-sig", errors="ignore")
        if text.count("�") > 3:
            text = raw.decode("latin-1")
        delim = ";" if text.split("\n")[0].count(";") >= text.split("\n")[0].count(",") else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    import unicodedata
    norm = lambda s: "".join(c for c in unicodedata.normalize("NFD", str(s).lower()) if unicodedata.category(c) != "Mn").strip()
    head = [norm(h) for h in rows[0]]

    def col(*keys):
        for i, h in enumerate(head):
            if any(k in h for k in keys):
                return i
        return None
    idx = {"amount": col("importe"), "charge": col("fecha cobro"), "name": col("nombre"), "nif": col("identificador", "nif"),
           "iban": col("iban"), "mandate": col("mandato") if col("mandato") != col("fecha mandato") else None,
           "mandate_date": col("fecha mandato"), "reference": col("referencia"), "concept": col("concepto"), "email": col("email", "correo"),
           "iva": None, "exempt": col("exento", "exenta")}
    for i, h in enumerate(head):
        if h == "mandato":
            idx["mandate"] = i
        if h in ("iva", "iva %", "% iva", "tipo iva", "tipo de iva", "iva (%)"):
            idx["iva"] = i
    out = []
    for r in rows[1:]:
        g = lambda k: (str(r[idx[k]]).strip() if idx[k] is not None and idx[k] < len(r) else "")
        ex = norm(g("exempt")) in ("si", "s", "x", "1", "true", "exento", "exenta")
        iva_raw = g("iva").replace("%", "")
        out.append({"amount": _num(g("amount")), "charge_date": g("charge")[:10], "name": g("name"), "nif": g("nif").upper(),
                    "iban": g("iban").replace(" ", "").upper(), "mandate": g("mandate"), "mandate_date": g("mandate_date")[:10],
                    "reference": g("reference"), "concept": g("concept"), "email": g("email"),
                    "iva_rate": 0.0 if ex else (_num(iva_raw) if iva_raw else 21.0), "vat_exempt": ex})
    return [x for x in out if x["iban"] or x["amount"]]


@rem.post("/importar")
async def import_remesa(file: UploadFile = File(...), name: str = Query(""), save_debtors: bool = Query(True), g=Depends(require_gestoria)):
    raw = await file.read()
    try:
        rows = _rows_from_file(file.filename or "", raw)
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo. Usa CSV (separado por ;) o Excel.")
    if not rows:
        raise HTTPException(status_code=400, detail="El archivo no tiene líneas de cobro.")
    lines, new_debtors = [], 0
    for x in rows:
        did = ""
        if save_debtors and x["iban"]:
            d = await db.remesa_debtors.find_one({"gestoria_id": g["id"], "iban": x["iban"], "mandate": x["mandate"]})
            if d:
                did = d["id"]
                if d.get("vat_exempt") and not x["vat_exempt"]:
                    x.update(vat_exempt=True, iva_rate=0.0)
                x["exempt_reason"] = d.get("exempt_reason", "")
            else:
                did = str(uuid.uuid4())
                await db.remesa_debtors.insert_one({**DebtorInput(name=x["name"], nif=x["nif"], iban=x["iban"], mandate=x["mandate"],
                                                                  mandate_date=x["mandate_date"], email=x["email"], amount=x["amount"],
                                                                  iva_rate=x["iva_rate"], vat_exempt=x["vat_exempt"],
                                                                  concept=x["concept"], reference=x["reference"]).model_dump(),
                                                    "id": did, "gestoria_id": g["id"], "created_at": _now()})
                new_debtors += 1
        lines.append({**LineInput(**{k: v for k, v in x.items() if k != "charge_date"}, debtor_id=did).model_dump(), "id": str(uuid.uuid4())})
    charge = next((x["charge_date"] for x in rows if x["charge_date"]), "")
    doc = {"id": str(uuid.uuid4()), "gestoria_id": g["id"], "name": name or (file.filename or "Remesa importada").rsplit(".", 1)[0],
           "charge_date": charge, "reference": "", "status": "borrador", "lines": lines, "created_at": _now()}
    await db.remesas.insert_one(doc)
    doc.pop("_id", None)
    return {"remesa": _summary(doc), "new_debtors": new_debtors}


@rem.get("/{rid}/exportar")
async def export_remesa(rid: str, formato: str = "xlsx", g=Depends(require_gestoria)):
    r = _summary(await _get(rid, g))
    rows = [[i, _r(ln.get("amount")), r.get("charge_date", ""), ln.get("name", ""), ln.get("nif", ""), ln.get("iban", ""), ln.get("mandate", ""),
             ln.get("mandate_date", ""), ln.get("reference", ""), ln.get("concept", ""), ln["base"], _rate(ln), ln["iva"],
             "Sí" if ln.get("vat_exempt") else "No", ln.get("invoice_number", "")] for i, ln in enumerate(r["lines"], 1)]
    safe = "".join(c if c.isalnum() else "_" for c in r["name"])[:60] or "remesa"
    if formato == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow(COLUMNS)
        w.writerows(rows)
        return Response(("\ufeff" + buf.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{safe}.csv"'})
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Remesa"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DBEAFE")
    for row in rows:
        ws.append(row)
        for c in (2, 11, 13):
            ws.cell(ws.max_row, c).number_format = "#,##0.00"
    for col, wd in zip("ABCDEFGHIJKLMNO", (6, 12, 13, 32, 18, 30, 18, 14, 34, 50, 14, 8, 12, 8, 16)):
        ws.column_dimensions[col].width = wd
    s = wb.create_sheet("Resumen")
    tot = _r(sum(x[1] for x in rows))
    for k, v in (("Remesa", r["name"]), ("Fecha de cobro", r.get("charge_date", "")), ("Referencia", r.get("reference", "")),
                 ("Nº de recibos", len(rows)), ("Base imponible", r["base_total"]), ("IVA", r["iva_total"]), ("Importe total", tot),
                 ("Pendiente de cobro", r["pending_total"]), ("Generada", datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M"))):
        s.append([k, v])
    s.column_dimensions["A"].width, s.column_dimensions["B"].width = 18, 40
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{safe}.xlsx"'})


# ---------- Detección de cobros (conciliación bancaria) ----------
def _label_hit(label: str, ln: dict) -> bool:
    lab = (label or "").upper()
    keys = [ln.get("mandate", ""), ln.get("nif", ""), ln.get("reference", "")] + [w for w in (ln.get("name") or "").upper().split() if len(w) > 3]
    return any(k and k.upper() in lab for k in keys)


async def detect_remesa_payments(gestoria_id: str, company_id: str) -> list:
    srv = __import__("server")
    found = []
    remesas = await db.remesas.find({"gestoria_id": gestoria_id, "status": {"$in": OPEN}}, {"_id": 0}).sort("created_at", -1).to_list(500)
    if not remesas:
        return found
    # Facturas de remesa cobradas desde la conciliación general → marcar el recibo como cobrado
    for r in remesas:
        ids = [ln["invoice_id"] for ln in r["lines"] if ln.get("invoice_id") and not ln.get("paid")]
        if not ids:
            continue
        paid = {i["id"]: i async for i in db.invoices.find({"id": {"$in": ids}, "status": "paid"}, {"_id": 0, "id": 1, "payment": 1})}
        for ln in r["lines"]:
            if ln.get("invoice_id") in paid:
                ln["paid"], ln["paid_date"] = True, ((paid[ln["invoice_id"]].get("payment") or {}).get("date") or datetime.now(timezone.utc).date().isoformat())[:10]
                found.append({"remesa": r["name"], "amount": _r(ln["amount"]), "detail": f"recibo de {ln.get('name')} (factura ya cobrada)"})
        if paid:
            r["status"] = await _save_lines(r)
    remesas = [r for r in remesas if r["status"] != "cobrada"]
    txs = await db.bank_transactions.find({"company_id": company_id, "matched": {"$ne": True}, "deleted": {"$ne": True}, "value": {"$gt": 0}}, {"_id": 0}).sort("date", -1).to_list(5000)
    for tx in txs:
        val = _r(tx.get("value"))
        hit = None
        for r in remesas:
            pend = [ln for ln in r["lines"] if not ln.get("paid")]
            if pend and abs(_r(sum(_r(ln["amount"]) for ln in pend)) - val) < 0.01:
                for ln in pend:
                    await _mark_line(r, ln, True, tx.get("date"), "bank")
                hit = (r, f"remesa completa ({len(pend)} recibos)")
                break
            cands = [ln for ln in pend if abs(_r(ln["amount"]) - val) < 0.01]
            ln = next((c for c in cands if _label_hit(tx.get("label", ""), c)), None) or (cands[0] if len(cands) == 1 and _label_hit(tx.get("label", ""), cands[0]) else None)
            if ln:
                await _mark_line(r, ln, True, tx.get("date"), "bank")
                hit = (r, f"recibo de {ln.get('name') or ln.get('iban')}")
                break
        if not hit:
            continue
        r, what = hit
        await _save_lines(r)
        await db.bank_transactions.update_one({"id": tx["id"]}, {"$set": {"matched": True, "matched_type": "remesa", "matched_id": r["id"]}})
        g = await db.users.find_one({"_id": ObjectId(gestoria_id)}, {"email": 1})
        msg = f"Cobro SEPA de {val:.2f} € del {tx.get('date', '')} conciliado con la remesa «{r['name']}»: {what}. Las facturas se han marcado como cobradas."
        await srv._notify(gestoria_id, company_id, "conciliacion", "Cobro de remesa detectado", msg, (g or {}).get("email"))
        found.append({"remesa": r["name"], "amount": val, "detail": what})
    return found


@rem.post("/comprobar-cobros")
async def check_payments(g=Depends(require_gestoria)):
    srv = __import__("server")
    cid = await srv.active_cid(g)
    return {"found": await detect_remesa_payments(g["id"], cid)}
