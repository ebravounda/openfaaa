"""Contabilidad (PGC): diario, mayor, sumas y saldos, balance, PyG, volumen, amortizaciones, asientos y categorías."""
import io
import uuid
from datetime import datetime, timezone, date, timedelta
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel

import pgc
from database import db
from auth import get_current_user

acc = APIRouter(prefix="/api/contabilidad", tags=["contabilidad"])


def _srv():
    import server
    return server


async def _company(user):
    return await _srv().active_company(user)


def _r(x):
    return round(float(x or 0) + 0.0, 2)


def _line(account, amount, side):
    """Línea de asiento; los importes negativos (rectificativas) pasan al lado contrario."""
    a = _r(amount)
    if not a:
        return None
    if a < 0:
        side, a = ("haber" if side == "debe" else "debe"), -a
    return {"account": account, "debe": a if side == "debe" else 0.0, "haber": a if side == "haber" else 0.0}


def _entry(d, concept, source, source_id, lines):
    lines = [x for x in lines if x]
    return {"date": (d or "")[:10], "concept": concept, "source": source, "source_id": source_id, "lines": lines} if lines else None


def _month_end(period: str) -> str:
    try:
        y, m = int(period[:4]), int(period[5:7])
        nxt = date(y + (m == 12), 1 if m == 12 else m + 1, 1)
        return (nxt - timedelta(days=1)).isoformat()
    except Exception:
        return period[:10]


def asset_amort(asset: dict, year: int) -> float:
    """Amortización lineal del año (por meses completos desde la puesta en funcionamiento)."""
    try:
        start = date.fromisoformat(asset["start_date"][:10])
    except Exception:
        return 0.0
    base = max(0.0, float(asset.get("cost", 0)) - float(asset.get("residual", 0) or 0))
    rate = float(asset.get("rate", 0) or 0)
    if not base or not rate or year < start.year:
        return 0.0
    monthly = base * rate / 100 / 12
    today = date.today()
    end_m = 12 if year < today.year else today.month
    if year > today.year:
        return 0.0
    first_m = start.month if year == start.year else 1
    months_before = max(0, (year - start.year) * 12 - (start.month - 1))
    done = monthly * months_before
    months = max(0, end_m - first_m + 1)
    return _r(max(0.0, min(monthly * months, base - done)))


async def build_entries(user, company, year: int) -> list:
    cid, uid = company["id"], user["id"]
    y = str(year)
    q = {"user_id": uid, "company_id": cid}
    out = []
    async for inv in db.invoices.find({**q, "issue_date": {"$regex": f"^{y}"}, "status": {"$ne": "anulada"}, "offsets_annulled": {"$ne": True}}, {"_id": 0}):
        cl = (inv.get("client") or {}).get("name", "")
        inc = pgc.INCOME_CATEGORIES.get(inv.get("income_category") or "", "705")
        e = _entry(inv["issue_date"], f"Factura {inv.get('number')} · {cl}", "factura", inv["id"], [
            _line("430", inv.get("total", 0), "debe"),
            _line("473", inv.get("irpf_amount", 0), "debe"),
            _line(inc, inv.get("base", 0), "haber"),
            _line("477", _r(inv.get("iva_amount", 0)) + _r(inv.get("re_amount", 0)), "haber"),
            _line("555", inv.get("suplidos_total", 0), "haber"),
        ])
        if e:
            out.append(e)
        if inv.get("status") == "paid":
            pdate = ((inv.get("payment") or {}).get("date") or inv.get("paid_at") or inv["issue_date"])[:10]
            if pdate[:4] == y:
                out.append(_entry(pdate, f"Cobro factura {inv.get('number')} · {cl}", "cobro", inv["id"],
                                  [_line("572", inv.get("total", 0), "debe"), _line("430", inv.get("total", 0), "haber")]))
    async for ex in db.expenses.find({**q, "date": {"$regex": f"^{y}"}}, {"_id": 0}):
        acct = pgc.EXPENSE_CATEGORIES.get(ex.get("category") or "General", "629")
        cred = "400" if acct.startswith("60") else "410"
        name = ex.get("vendor_name", "")
        e = _entry(ex["date"], f"Gasto {ex.get('invoice_number') or ''} · {name}".replace("  ", " "), "gasto", ex["id"], [
            _line(acct, ex.get("base", 0), "debe"), _line("472", ex.get("iva_amount", 0), "debe"),
            _line(cred, ex.get("total", 0), "haber")])
        if e:
            out.append(e)
        if ex.get("reconciled"):
            out.append(_entry(ex["date"], f"Pago gasto · {name}", "pago", ex["id"],
                              [_line(cred, ex.get("total", 0), "debe"), _line("572", ex.get("total", 0), "haber")]))
    async for p in db.payslips.find({**q, "period": {"$regex": f"^{y}"}}, {"_id": 0}):
        who = p.get("employee_name") or "trabajador"
        out.append(_entry(_month_end(p.get("period", "")), f"Nómina {p.get('period')} · {who}", "nomina", p.get("id"), [
            _line("640", p.get("gross", 0), "debe"), _line("642", p.get("company_ss", 0), "debe"),
            _line("476", _r(p.get("ss_amount", 0)) + _r(p.get("company_ss", 0)), "haber"),
            _line("4751", p.get("irpf_amount", 0), "haber"), _line("460", p.get("other_deductions", 0), "haber"),
            _line("465", p.get("net", 0), "haber")]))
    async for a in db.fixed_assets.find({"company_id": cid}, {"_id": 0}):
        acct = a.get("account", "219")
        if a.get("register_purchase") and (a.get("start_date") or "")[:4] == y:
            out.append(_entry(a["start_date"], f"Alta inmovilizado · {a.get('name')}", "inmovilizado", a["id"], [
                _line(acct, a.get("cost", 0), "debe"), _line("472", a.get("iva_amount", 0), "debe"),
                _line("523", _r(a.get("cost", 0)) + _r(a.get("iva_amount", 0)), "haber")]))
        amt = asset_amort(a, year)
        if amt:
            today = date.today()
            d = f"{y}-12-31" if year < today.year else _month_end(today.isoformat()[:7])
            exp_acc = "680" if acct.startswith("20") else "681"
            amort_acc = ("280" if acct.startswith("20") else "281") + acct[2:3]
            out.append(_entry(d, f"Amortización {y} · {a.get('name')}", "amortizacion", a["id"],
                              [_line(exp_acc, amt, "debe"), _line(amort_acc, amt, "haber")]))
    async for m in db.journal_entries.find({"company_id": cid, "date": {"$regex": f"^{y}"}}, {"_id": 0}):
        out.append({"date": m["date"], "concept": m.get("concept", ""), "source": "manual", "source_id": m["id"], "lines": m.get("lines", [])})
    out = [e for e in out if e]
    out.sort(key=lambda e: (e["date"], e["source"] != "manual", e["concept"]))
    for i, e in enumerate(out, 1):
        e["number"] = i
        e["debe"] = _r(sum(x["debe"] for x in e["lines"]))
        e["haber"] = _r(sum(x["haber"] for x in e["lines"]))
    return out


def _balances(entries) -> dict:
    bal = {}
    for e in entries:
        for ln in e["lines"]:
            b = bal.setdefault(ln["account"], {"debe": 0.0, "haber": 0.0})
            b["debe"] += ln["debe"]
            b["haber"] += ln["haber"]
    return bal


async def _ctx(user, year):
    company = await _company(user)
    plan = company.get("pgc_plan") or "pymes"
    year = year or datetime.now(timezone.utc).year
    return company, plan, year, await build_entries(user, company, year)


def _sumas(entries, plan):
    rows = []
    for code, b in sorted(_balances(entries).items()):
        s = _r(b["debe"] - b["haber"])
        rows.append({"account": code, "name": pgc.account_name(code, plan), "debe": _r(b["debe"]), "haber": _r(b["haber"]),
                     "saldo_deudor": s if s > 0 else 0.0, "saldo_acreedor": -s if s < 0 else 0.0})
    return rows


def _pyg(entries):
    bal = _balances(entries)
    groups = [(lbl, pre) for lbl, pre in pgc.PYG if not lbl.startswith("__")]
    vals = {lbl: 0.0 for lbl, _ in groups}
    detail = {lbl: [] for lbl, _ in groups}
    for code, b in bal.items():
        if code[0] not in "67":
            continue
        best = max(groups, key=lambda g: pgc.match_prefix(code, g[1]))
        lbl = best[0] if pgc.match_prefix(code, best[1]) else ("7. Otros gastos de explotación" if code[0] == "6" else "5. Otros ingresos de explotación")
        v = b["haber"] - b["debe"]
        vals[lbl] += v
        detail[lbl].append({"account": code, "name": pgc.account_name(code), "amount": _r(v)})
    rows, A, B = [], 0.0, 0.0
    for lbl, pre in pgc.PYG:
        if lbl == "__A__":
            A = sum(vals[g[0]] for g in groups[:12]); rows.append({"label": pre, "amount": _r(A), "total": True})
        elif lbl == "__B__":
            B = sum(vals[g[0]] for g in groups[12:17]); rows.append({"label": pre, "amount": _r(B), "total": True})
        elif lbl == "__C__":
            rows.append({"label": pre, "amount": _r(A + B), "total": True})
        elif lbl == "__D__":
            rows.append({"label": pre, "amount": _r(A + B + vals["18. Impuestos sobre beneficios"]), "total": True})
        else:
            rows.append({"label": lbl, "amount": _r(vals[lbl]), "detail": detail[lbl]})
    return rows, _r(A + B + vals["18. Impuestos sobre beneficios"])


def _balance(entries, plan):
    bal = _balances(entries)
    _, result = _pyg(entries)
    sides = {"activo": pgc.BALANCE_ACTIVO, "pasivo": pgc.BALANCE_PASIVO}
    caps = [(side, sec, lbl, pre) for side, secs in sides.items() for sec, items in secs for lbl, pre in items]
    vals = {(c[1], c[2]): [] for c in caps}
    for code, b in bal.items():
        if code[0] in "67":
            continue
        best = max(caps, key=lambda c: pgc.match_prefix(code, c[3]))
        if not pgc.match_prefix(code, best[3]):
            s = b["debe"] - b["haber"]
            best = next(c for c in caps if c[2] == ("III. Deudores comerciales y otras cuentas a cobrar" if s >= 0 else "II. Deudas a corto plazo"))
        v = (b["debe"] - b["haber"]) if best[0] == "activo" else (b["haber"] - b["debe"])
        vals[(best[1], best[2])].append({"account": code, "name": pgc.account_name(code, plan), "amount": _r(v)})
    out = {}
    for side, secs in sides.items():
        sections, side_total = [], 0.0
        for sec, items in secs:
            lines = []
            for lbl, pre in items:
                if plan == "pymes" and lbl in pgc.PYMES_SKIP:
                    continue
                det = vals[(sec, lbl)]
                amt = result if pre == ["__RESULT__"] else sum(d["amount"] for d in det)
                lines.append({"label": lbl, "amount": _r(amt), "detail": det})
            st = _r(sum(x["amount"] for x in lines))
            side_total += st
            sections.append({"label": sec, "amount": st, "lines": lines})
        out[side] = {"sections": sections, "total": _r(side_total)}
    out["cuadra"] = abs(out["activo"]["total"] - out["pasivo"]["total"]) < 0.02
    out["modelo"] = "Modelo normal (PGC RD 1514/2007)" if plan == "normal" else "Modelo PYMES (PGC RD 1515/2007)"
    return out


@acc.get("/config")
async def get_config(user=Depends(get_current_user)):
    c = await _company(user)
    return {"plan": c.get("pgc_plan") or "pymes", "company": c.get("name", "")}


class ConfigInput(BaseModel):
    plan: str


@acc.put("/config")
async def set_config(data: ConfigInput, user=Depends(get_current_user)):
    if data.plan not in ("pymes", "normal"):
        raise HTTPException(status_code=400, detail="Plan no válido (pymes o normal)")
    c = await _company(user)
    await db.companies.update_one({"id": c["id"]}, {"$set": {"pgc_plan": data.plan}})
    return {"plan": data.plan}


@acc.get("/cuentas")
async def cuentas(user=Depends(get_current_user)):
    c = await _company(user)
    return pgc.catalogue(c.get("pgc_plan") or "pymes")


@acc.get("/diario")
async def diario(year: Optional[int] = None, user=Depends(get_current_user)):
    company, plan, year, entries = await _ctx(user, year)
    for e in entries:
        for ln in e["lines"]:
            ln["name"] = pgc.account_name(ln["account"], plan)
    return {"year": year, "entries": entries, "debe": _r(sum(e["debe"] for e in entries)), "haber": _r(sum(e["haber"] for e in entries))}


@acc.get("/mayor")
async def mayor(year: Optional[int] = None, account: str = "", user=Depends(get_current_user)):
    company, plan, year, entries = await _ctx(user, year)
    if not account:
        return {"year": year, "accounts": _sumas(entries, plan)}
    moves, saldo = [], 0.0
    for e in entries:
        for ln in e["lines"]:
            if ln["account"].startswith(account):
                saldo += ln["debe"] - ln["haber"]
                moves.append({"date": e["date"], "number": e["number"], "concept": e["concept"], "account": ln["account"],
                              "debe": ln["debe"], "haber": ln["haber"], "saldo": _r(saldo)})
    return {"year": year, "account": account, "name": pgc.account_name(account, plan), "moves": moves,
            "debe": _r(sum(m["debe"] for m in moves)), "haber": _r(sum(m["haber"] for m in moves)), "saldo": _r(saldo)}


@acc.get("/sumas-saldos")
async def sumas_saldos(year: Optional[int] = None, user=Depends(get_current_user)):
    company, plan, year, entries = await _ctx(user, year)
    rows = _sumas(entries, plan)
    tot = {k: _r(sum(r[k] for r in rows)) for k in ("debe", "haber", "saldo_deudor", "saldo_acreedor")}
    return {"year": year, "rows": rows, "totals": tot}


@acc.get("/balance")
async def balance(year: Optional[int] = None, user=Depends(get_current_user)):
    company, plan, year, entries = await _ctx(user, year)
    return {"year": year, "plan": plan, **_balance(entries, plan)}


@acc.get("/pyg")
async def pyg(year: Optional[int] = None, user=Depends(get_current_user)):
    company, plan, year, entries = await _ctx(user, year)
    rows, result = _pyg(entries)
    return {"year": year, "plan": plan, "rows": rows, "result": result}


@acc.get("/volumen")
async def volumen(year: Optional[int] = None, user=Depends(get_current_user)):
    company = await _company(user)
    year = year or datetime.now(timezone.utc).year
    q = {"user_id": user["id"], "company_id": company["id"]}
    quarters = [{"q": f"T{i}", "ventas": 0.0, "iva_repercutido": 0.0, "compras": 0.0, "iva_soportado": 0.0, "facturas": 0, "gastos": 0} for i in range(1, 5)]
    clients, vendors = {}, {}
    async for inv in db.invoices.find({**q, "issue_date": {"$regex": f"^{year}"}, "status": {"$ne": "anulada"}}, {"_id": 0}):
        t = quarters[(int(inv["issue_date"][5:7]) - 1) // 3]
        t["ventas"] += inv.get("base", 0) or 0; t["iva_repercutido"] += inv.get("iva_amount", 0) or 0; t["facturas"] += 1
        cl = inv.get("client") or {}
        k = cl.get("nif") or cl.get("name", "")
        c = clients.setdefault(k, {"name": cl.get("name", ""), "nif": cl.get("nif", ""), "base": 0.0, "total": 0.0, "count": 0})
        c["base"] += inv.get("base", 0) or 0; c["total"] += _r(inv.get("base", 0)) + _r(inv.get("iva_amount", 0)) + _r(inv.get("re_amount", 0)); c["count"] += 1
    async for ex in db.expenses.find({**q, "date": {"$regex": f"^{year}"}}, {"_id": 0}):
        t = quarters[(int(ex["date"][5:7]) - 1) // 3]
        t["compras"] += ex.get("base", 0) or 0; t["iva_soportado"] += ex.get("iva_amount", 0) or 0; t["gastos"] += 1
        k = ex.get("vendor_nif") or ex.get("vendor_name", "")
        v = vendors.setdefault(k, {"name": ex.get("vendor_name", ""), "nif": ex.get("vendor_nif", ""), "base": 0.0, "total": 0.0, "count": 0})
        v["base"] += ex.get("base", 0) or 0; v["total"] += ex.get("total", 0) or 0; v["count"] += 1
    for t in quarters:
        for k in ("ventas", "iva_repercutido", "compras", "iva_soportado"):
            t[k] = _r(t[k])

    def top(d):
        rows = sorted(d.values(), key=lambda x: -x["total"])
        for r in rows:
            r["base"], r["total"] = _r(r["base"]), _r(r["total"])
            r["modelo_347"] = r["total"] > 3005.06
        return rows
    ventas = _r(sum(t["ventas"] for t in quarters))
    return {"year": year, "quarters": quarters, "ventas": ventas, "compras": _r(sum(t["compras"] for t in quarters)),
            "clientes": top(clients), "proveedores": top(vendors)}


@acc.get("/trimestre")
async def trimestre(user=Depends(get_current_user)):
    today = date.today()
    qn = (today.month - 1) // 3 + 1
    start = date(today.year, 3 * qn - 2, 1)
    end = (date(today.year + (qn == 4), 1 if qn == 4 else 3 * qn + 1, 1) - timedelta(days=1))
    deadline = date(today.year + 1, 1, 30) if qn == 4 else date(today.year, 3 * qn + 1, 20)
    company = await _company(user)
    q = {"user_id": user["id"], "company_id": company["id"]}
    rng = {"$gte": start.isoformat(), "$lte": end.isoformat() + "T99"}
    ing = iva_r = gas = iva_s = 0.0
    n_inv = n_exp = 0
    async for inv in db.invoices.find({**q, "issue_date": rng, "status": {"$ne": "anulada"}}, {"base": 1, "iva_amount": 1}):
        ing += inv.get("base", 0) or 0; iva_r += inv.get("iva_amount", 0) or 0; n_inv += 1
    async for ex in db.expenses.find({**q, "date": rng}, {"base": 1, "iva_amount": 1}):
        gas += ex.get("base", 0) or 0; iva_s += ex.get("iva_amount", 0) or 0; n_exp += 1
    return {"quarter": qn, "label": f"{qn}T {today.year}", "year": today.year, "start": start.isoformat(), "end": end.isoformat(),
            "deadline": deadline.isoformat(), "days_left_quarter": (end - today).days, "days_to_deadline": (deadline - today).days,
            "progress": round(((today - start).days + 1) / ((end - start).days + 1) * 100),
            "ingresos": _r(ing), "gastos": _r(gas), "iva_repercutido": _r(iva_r), "iva_soportado": _r(iva_s),
            "iva_estimado": _r(iva_r - iva_s), "facturas": n_inv, "gastos_count": n_exp}


# ---------- Amortizaciones ----------
class AssetInput(BaseModel):
    name: str
    asset_type: str = "otro"
    account: str = ""
    cost: float
    iva_amount: float = 0.0
    residual: float = 0.0
    start_date: str
    rate: float = 0.0
    register_purchase: bool = False


def _asset_doc(data: AssetInput) -> dict:
    t = next((x for x in pgc.AMORT_TABLE if x["key"] == data.asset_type), pgc.AMORT_TABLE[-2])
    rate = data.rate or t["rate"]
    if rate > t["rate"]:
        raise HTTPException(status_code=400, detail=f"El coeficiente máximo fiscal para {t['label']} es {t['rate']} %.")
    if data.cost <= 0:
        raise HTTPException(status_code=400, detail="El valor de adquisición debe ser mayor que 0.")
    d = data.model_dump()
    d.update({"account": data.account or t["account"], "rate": rate, "type_label": t["label"]})
    return d


def _schedule(a: dict) -> list:
    rows, acum = [], 0.0
    base = float(a["cost"]) - float(a.get("residual") or 0)
    y0 = int(a["start_date"][:4])
    for y in range(y0, y0 + 120):
        start = date.fromisoformat(a["start_date"][:10])
        monthly = base * a["rate"] / 100 / 12
        months = 13 - start.month if y == y0 else 12
        amt = _r(min(monthly * months, base - acum))
        if amt <= 0:
            break
        acum = _r(acum + amt)
        rows.append({"year": y, "amortizacion": amt, "acumulada": acum, "valor_neto": _r(float(a["cost"]) - acum)})
    return rows


@acc.get("/tablas-amortizacion")
async def tablas(user=Depends(get_current_user)):
    return pgc.AMORT_TABLE


@acc.get("/activos")
async def list_assets(user=Depends(get_current_user)):
    c = await _company(user)
    year = date.today().year
    rows = await db.fixed_assets.find({"company_id": c["id"]}, {"_id": 0}).sort("start_date", -1).to_list(1000)
    for a in rows:
        sch = _schedule(a)
        a["amort_year"] = asset_amort(a, year)
        prev = sum(r["amortizacion"] for r in sch if r["year"] < year)
        a["acumulada"] = _r(prev + a["amort_year"])
        a["valor_neto"] = _r(a["cost"] - a["acumulada"])
        a["schedule"] = sch
    return rows


@acc.post("/activos")
async def create_asset(data: AssetInput, user=Depends(get_current_user)):
    c = await _company(user)
    doc = {**_asset_doc(data), "id": str(uuid.uuid4()), "company_id": c["id"], "user_id": user["id"],
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.fixed_assets.insert_one(doc)
    doc.pop("_id", None)
    return doc


@acc.put("/activos/{asset_id}")
async def update_asset(asset_id: str, data: AssetInput, user=Depends(get_current_user)):
    c = await _company(user)
    r = await db.fixed_assets.update_one({"id": asset_id, "company_id": c["id"]}, {"$set": _asset_doc(data)})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Activo no encontrado")
    return {"status": "ok"}


@acc.delete("/activos/{asset_id}")
async def delete_asset(asset_id: str, user=Depends(get_current_user)):
    c = await _company(user)
    await db.fixed_assets.delete_one({"id": asset_id, "company_id": c["id"]})
    return {"status": "ok"}


# ---------- Asientos manuales ----------
class JLine(BaseModel):
    account: str
    debe: float = 0.0
    haber: float = 0.0


class JournalInput(BaseModel):
    date: str
    concept: str
    lines: List[JLine]


def _journal_doc(data: JournalInput) -> dict:
    lines = [{"account": ln.account.strip(), "debe": _r(ln.debe), "haber": _r(ln.haber)} for ln in data.lines if ln.account.strip() and (ln.debe or ln.haber)]
    if len(lines) < 2:
        raise HTTPException(status_code=400, detail="El asiento necesita al menos dos líneas con cuenta e importe.")
    if any(not ln["account"].isdigit() for ln in lines):
        raise HTTPException(status_code=400, detail="Las cuentas deben ser numéricas (p. ej. 572, 430).")
    d, h = _r(sum(x["debe"] for x in lines)), _r(sum(x["haber"] for x in lines))
    if abs(d - h) > 0.005:
        raise HTTPException(status_code=400, detail=f"El asiento no cuadra: Debe {d:.2f} € ≠ Haber {h:.2f} €.")
    return {"date": data.date[:10], "concept": data.concept.strip(), "lines": lines}


@acc.get("/asientos")
async def list_journal(year: Optional[int] = None, user=Depends(get_current_user)):
    c = await _company(user)
    q = {"company_id": c["id"]}
    if year:
        q["date"] = {"$regex": f"^{year}"}
    return await db.journal_entries.find(q, {"_id": 0}).sort("date", -1).to_list(5000)


@acc.post("/asientos")
async def create_journal(data: JournalInput, user=Depends(get_current_user)):
    c = await _company(user)
    doc = {**_journal_doc(data), "id": str(uuid.uuid4()), "company_id": c["id"], "user_id": user["id"],
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.journal_entries.insert_one(doc)
    doc.pop("_id", None)
    return doc


@acc.put("/asientos/{entry_id}")
async def update_journal(entry_id: str, data: JournalInput, user=Depends(get_current_user)):
    c = await _company(user)
    r = await db.journal_entries.update_one({"id": entry_id, "company_id": c["id"]}, {"$set": _journal_doc(data)})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Asiento no encontrado")
    return {"status": "ok"}


@acc.delete("/asientos/{entry_id}")
async def delete_journal(entry_id: str, user=Depends(get_current_user)):
    c = await _company(user)
    await db.journal_entries.delete_one({"id": entry_id, "company_id": c["id"]})
    return {"status": "ok"}


# ---------- Categorías ----------
@acc.get("/categorias")
async def categorias(year: Optional[int] = None, user=Depends(get_current_user)):
    c = await _company(user)
    q = {"user_id": user["id"], "company_id": c["id"]}
    if year:
        q_e, q_i = {**q, "date": {"$regex": f"^{year}"}}, {**q, "issue_date": {"$regex": f"^{year}"}}
    else:
        q_e, q_i = q, q
    exps = await db.expenses.find(q_e, {"_id": 0, "id": 1, "date": 1, "vendor_name": 1, "description": 1, "category": 1, "base": 1}).sort("date", -1).to_list(3000)
    invs = await db.invoices.find({**q_i, "status": {"$ne": "anulada"}}, {"_id": 0, "id": 1, "number": 1, "issue_date": 1, "client": 1, "income_category": 1, "base": 1, "line_items": 1}).sort("issue_date", -1).to_list(3000)
    for e in exps:
        e["account"] = pgc.EXPENSE_CATEGORIES.get(e.get("category") or "General", "629")
        e["suggested"] = pgc.guess_category(f"{e.get('vendor_name', '')} {e.get('description', '')}", pgc.EXPENSE_KEYWORDS)
    for i in invs:
        i["income_category"] = i.get("income_category") or "Prestación de servicios"
        i["account"] = pgc.INCOME_CATEGORIES.get(i["income_category"], "705")
        i["client_name"] = (i.pop("client", None) or {}).get("name", "")
        text = " ".join(f"{li.get('description', '')} {li.get('detail', '')}" for li in (i.pop("line_items", None) or []))
        i["suggested"] = pgc.guess_category(text, pgc.INCOME_KEYWORDS)
    totals = {}
    for e in exps:
        k = e.get("category") or "General"
        totals[k] = _r(totals.get(k, 0) + (e.get("base") or 0))
    return {"expense_categories": [{"name": k, "account": v, "account_name": pgc.account_name(v)} for k, v in pgc.EXPENSE_CATEGORIES.items()],
            "income_categories": [{"name": k, "account": v, "account_name": pgc.account_name(v)} for k, v in pgc.INCOME_CATEGORIES.items()],
            "expenses": exps, "invoices": invs, "expense_totals": totals}


class CategoryInput(BaseModel):
    kind: str
    id: str
    category: str


@acc.put("/categorias")
async def set_category(data: CategoryInput, user=Depends(get_current_user)):
    c = await _company(user)
    q = {"id": data.id, "user_id": user["id"], "company_id": c["id"]}
    if data.kind == "gasto" and data.category in pgc.EXPENSE_CATEGORIES:
        r = await db.expenses.update_one(q, {"$set": {"category": data.category}})
    elif data.kind == "factura" and data.category in pgc.INCOME_CATEGORIES:
        r = await db.invoices.update_one(q, {"$set": {"income_category": data.category}})
    else:
        raise HTTPException(status_code=400, detail="Categoría no válida")
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Documento no encontrado")
    return {"status": "ok"}


@acc.post("/auto-categorizar")
async def auto_categorize(user=Depends(get_current_user)):
    c = await _company(user)
    q = {"user_id": user["id"], "company_id": c["id"]}
    n_e = n_i = 0
    async for e in db.expenses.find({**q, "category": {"$in": ["General", "Otros", "", None]}}, {"_id": 0, "id": 1, "vendor_name": 1, "description": 1}):
        g = pgc.guess_category(f"{e.get('vendor_name', '')} {e.get('description', '')}", pgc.EXPENSE_KEYWORDS)
        if g:
            await db.expenses.update_one({"id": e["id"], **q}, {"$set": {"category": g}}); n_e += 1
    async for i in db.invoices.find({**q, "income_category": {"$in": ["", None]}}, {"_id": 0, "id": 1, "line_items": 1}):
        text = " ".join(f"{li.get('description', '')} {li.get('detail', '')}" for li in (i.get("line_items") or []))
        g = pgc.guess_category(text, pgc.INCOME_KEYWORDS) or "Prestación de servicios"
        await db.invoices.update_one({"id": i["id"], **q}, {"$set": {"income_category": g}}); n_i += 1
    return {"gastos": n_e, "facturas": n_i}


# ---------- Exportación Excel ----------
@acc.get("/exportar")
async def exportar(libro: str = Query(...), year: Optional[int] = None, user=Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    company, plan, year, entries = await _ctx(user, year)
    wb = Workbook(); ws = wb.active
    bold, head = Font(bold=True), PatternFill("solid", fgColor="DBEAFE")
    ws.append([f"{company.get('name', '')} · NIF {company.get('nif', '')}"]); ws["A1"].font = Font(bold=True, size=13)

    def header(cols):
        ws.append(cols)
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = bold, head
    if libro == "diario":
        ws.title = "Libro diario"; ws.append([f"Libro diario · Ejercicio {year}"]); header(["Asiento", "Fecha", "Cuenta", "Descripción cuenta", "Concepto", "Debe", "Haber"])
        for e in entries:
            for ln in e["lines"]:
                ws.append([e["number"], e["date"], ln["account"], pgc.account_name(ln["account"], plan), e["concept"], ln["debe"] or None, ln["haber"] or None])
    elif libro == "mayor":
        ws.title = "Libro mayor"; ws.append([f"Libro mayor · Ejercicio {year}"])
        for code in sorted(_balances(entries)):
            ws.append([]); ws.append([f"{code} · {pgc.account_name(code, plan)}"]); ws[ws.max_row][0].font = bold
            header(["Fecha", "Asiento", "Concepto", "Debe", "Haber", "Saldo"]); s = 0.0
            for e in entries:
                for ln in e["lines"]:
                    if ln["account"] == code:
                        s += ln["debe"] - ln["haber"]
                        ws.append([e["date"], e["number"], e["concept"], ln["debe"] or None, ln["haber"] or None, _r(s)])
    elif libro == "sumas":
        ws.title = "Sumas y saldos"; ws.append([f"Balance de sumas y saldos · {year}"]); header(["Cuenta", "Descripción", "Debe", "Haber", "Saldo deudor", "Saldo acreedor"])
        for r in _sumas(entries, plan):
            ws.append([r["account"], r["name"], r["debe"], r["haber"], r["saldo_deudor"], r["saldo_acreedor"]])
    elif libro == "balance":
        b = _balance(entries, plan); ws.title = "Balance de situación"; ws.append([f"Balance de situación · 31/12/{year} · {b['modelo']}"])
        for side, title in (("activo", "ACTIVO"), ("pasivo", "PATRIMONIO NETO Y PASIVO")):
            ws.append([]); header([title, "Importe"])
            for sec in b[side]["sections"]:
                ws.append([sec["label"], sec["amount"]]); ws[ws.max_row][0].font = bold
                for ln in sec["lines"]:
                    ws.append(["    " + ln["label"], ln["amount"]])
            ws.append([f"TOTAL {title}", b[side]["total"]]); ws[ws.max_row][0].font = bold
    elif libro == "pyg":
        rows, _ = _pyg(entries); ws.title = "Pérdidas y ganancias"; ws.append([f"Cuenta de pérdidas y ganancias · {year}"]); header(["Concepto", "Importe"])
        for r in rows:
            ws.append([r["label"], r["amount"]])
            if r.get("total"):
                ws[ws.max_row][0].font = bold
    else:
        raise HTTPException(status_code=400, detail="Libro no válido")
    for col, w in zip("ABCDEFG", (14, 14, 40, 40, 50, 14, 14)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{libro}_{year}.xlsx"'})
